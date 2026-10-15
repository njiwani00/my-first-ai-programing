"""
STEP 1 - LOAD: read documents from a folder.

Supported formats: .txt, .md, .pdf, .csv and .docx
Each file becomes one Document with its text and where it came from (source).
Remembering the source is what lets us show citations later.

Key idea: RAG works on TEXT. Every file type is simply converted to plain text
here, and the rest of the pipeline never needs to know what the file was.

  .txt / .md  -> read as they are
  .pdf        -> extract the text from each page (scanned image PDFs have no text!)
  .docx       -> join the paragraphs of the Word document
  .csv        -> turn every ROW into a readable sentence, for example
                 "code: CS201 | course: Data Structures | instructor: Dr. Meera Rao"
                 A row of bare numbers means nothing on its own; adding the
                 column names gives each row the context search needs.
"""

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import List

SUPPORTED = {".txt", ".md", ".pdf", ".csv", ".docx"}


@dataclass
class Document:
    text: str
    source: str  # file name, used for citations


def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError("Reading PDFs needs pypdf:  pip install pypdf") from exc
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def read_docx(path: Path) -> str:
    try:
        import docx  # the python-docx package
    except ImportError as exc:
        raise ImportError("Reading Word files needs python-docx:  pip install python-docx") from exc
    document = docx.Document(str(path))
    parts = []
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "") if p.style is not None else ""
        if style.startswith("Heading") or style == "Title":
            # Mark Word headings like Markdown ("## Heading") so the chunker
            # keeps them on their own line instead of gluing them to the next sentence.
            level = style.replace("Heading", "").strip()
            hashes = "#" * (int(level) if level.isdigit() else 1)
            parts.append(f"{hashes} {text}")
        else:
            parts.append(text)
    # Tables in Word files: one line per row, cells separated by " | "
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells) + ".")
    return "\n\n".join(parts)


def read_csv(path: Path) -> str:
    """Each row becomes one line of 'column: value' pairs, separated by blank lines."""
    with open(path, encoding="utf-8-sig", errors="ignore", newline="") as f:
        reader = csv.DictReader(f)
        rows = []
        for row in reader:
            parts = [f"{col.strip()}: {val.strip()}" for col, val in row.items() if col and val and val.strip()]
            if parts:
                rows.append(" | ".join(parts) + ".")
    # Blank lines between rows let the chunker treat each row as one unit,
    # so a row is never cut in half.
    return "\n\n".join(rows)


def load_file(path: Path) -> Document:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = read_pdf(path)
    elif suffix == ".docx":
        text = read_docx(path)
    elif suffix == ".csv":
        text = read_csv(path)
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")
    return Document(text=text.strip(), source=path.name)


def load_documents(folder: str) -> List[Document]:
    """Load every supported file in `folder` (including sub-folders)."""
    folder_path = Path(folder)
    if not folder_path.exists():
        raise FileNotFoundError(f"Data folder not found: {folder_path}")

    documents = []
    for path in sorted(folder_path.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED and not path.name.startswith("~$"):
            doc = load_file(path)
            if doc.text:  # skip empty files
                documents.append(doc)
    return documents
